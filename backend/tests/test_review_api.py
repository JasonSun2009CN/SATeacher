"""Batch-3 review APIs: hand-written explanations, vocabulary grid + xlsx, AI answer."""

from __future__ import annotations

from io import BytesIO

from app import llm
from tests.conftest import MINI_SATMD, upload

AI_BANK = """---
satmd: 1
title: "AI Bank"
source: "ai.md"
lang: en
answers: external
---

:::q {#Q001 sec=rw no=1}
@material
The passage describes a scholar who counts bees among the orchards.

@stem
Which choice best completes the sentence? The scholar found the data ___ .

- A. one
- B. two
- C. three
- D. four

! answer: B
! explain: My private notes about why B is right.
:::
"""


def _mini(client) -> tuple[int, list[dict]]:
    resp = upload(client, "mini.sat.md", MINI_SATMD)
    assert resp.status_code == 200, resp.text
    doc_id = resp.json()["id"]
    questions = client.get(f"/api/documents/{doc_id}/questions").json()
    return doc_id, questions


# --------------------------------------------------------------------------
# explanations
# --------------------------------------------------------------------------


def test_explain_save_and_clear(client) -> None:
    doc_id, questions = _mini(client)
    qid = questions[0]["id"]

    ok = client.put(
        f"/api/documents/{doc_id}/questions/{qid}/explain",
        json={"content": "Substitute x = 5 and simplify."},
    )
    assert ok.status_code == 200, ok.text
    assert ok.json()["explain"] == "Substitute x = 5 and simplify."
    assert client.get(f"/api/documents/{doc_id}/questions").json()[0]["explain"] == (
        "Substitute x = 5 and simplify."
    )

    # blank text clears it back to null
    cleared = client.put(
        f"/api/documents/{doc_id}/questions/{qid}/explain", json={"content": "   "}
    )
    assert cleared.json()["explain"] is None
    assert client.get(f"/api/documents/{doc_id}/questions").json()[0]["explain"] is None


def test_explain_unknown_question_404(client) -> None:
    doc_id, _questions = _mini(client)
    assert (
        client.put(
            f"/api/documents/{doc_id}/questions/999999/explain",
            json={"content": "x"},
        ).status_code
        == 404
    )
    assert (
        client.put("/api/documents/999999/questions/1/explain", json={"content": "x"}).status_code
        == 404
    )


# --------------------------------------------------------------------------
# vocabulary grid
# --------------------------------------------------------------------------


def test_words_default_then_roundtrip(client) -> None:
    doc_id, _questions = _mini(client)

    fresh = client.get(f"/api/documents/{doc_id}/words")
    assert fresh.status_code == 200
    body = fresh.json()
    assert body["version"] == 2
    assert [c["name"] for c in body["columns"]] == ["Word", "Meaning", "Notes"]
    assert all(c["id"] and c["width"] >= 60 for c in body["columns"])
    assert body["rows"] == []
    assert body["view"] == {"sort": None, "filter": {}}

    saved = client.put(
        f"/api/documents/{doc_id}/words",
        json={
            "version": 2,
            "columns": [
                {"id": "c1", "name": "Word", "width": 180},
                {"id": "c2", "name": "Meaning", "width": 260},
                {"id": "c3", "name": "Notes", "width": 220},
                {"id": "c4", "name": "Tag", "width": 120},
            ],
            "rows": [
                {"id": "r1", "cells": {"c1": "elate", "c2": "make happy", "c3": "verb", "c4": "rw"}},
                {"id": "r2", "cells": {"c1": "cadence", "c2": "rhythm"}},
            ],
        },
    )
    assert saved.status_code == 200, saved.text
    # missing cells are filled with empty strings, keyed by stable column id
    assert saved.json()["rows"][1]["cells"]["c3"] == ""

    again = client.get(f"/api/documents/{doc_id}/words").json()
    assert [c["name"] for c in again["columns"]] == ["Word", "Meaning", "Notes", "Tag"]
    assert again["rows"][0]["cells"]["c1"] == "elate"
    assert again["rows"][0]["cells"]["c4"] == "rw"


def test_words_validation(client) -> None:
    doc_id, _questions = _mini(client)
    base = f"/api/documents/{doc_id}/words"

    assert client.put(base, json={"columns": [], "rows": []}).status_code == 400
    assert (
        client.put(
            base,
            json={"columns": [{"name": f"c{i}"} for i in range(13)], "rows": []},
        ).status_code
        == 400
    )
    assert (
        client.put(
            base,
            json={"columns": [{"name": "Word"}], "rows": [{"cells": {"c1": "x"}}] * 501},
        ).status_code
        == 400
    )
    # a blank column name is rejected
    assert client.put(base, json={"columns": [{"name": "   "}], "rows": []}).status_code == 400
    assert (
        client.put(
            "/api/documents/999999/words", json={"columns": [{"name": "Word"}], "rows": []}
        ).status_code
        == 404
    )


def test_words_xlsx_export(client) -> None:
    doc_id, _questions = _mini(client)
    client.put(
        f"/api/documents/{doc_id}/words",
        json={
            "columns": [
                {"id": "c1", "name": "Word", "width": 180},
                {"id": "c2", "name": "Meaning", "width": 240},
            ],
            "rows": [
                {"id": "r1", "cells": {"c1": "ephemeral", "c2": "short-lived"}},
                {"id": "r2", "cells": {"c1": "sturdy", "c2": "strong"}},
            ],
        },
    )

    resp = client.get(f"/api/documents/{doc_id}/words/export")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "vocabulary.xlsx" in resp.headers["content-disposition"]
    assert resp.content[:2] == b"PK"  # zip container

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(resp.content))
    ws = wb.active
    assert ws.title == "Vocabulary"
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert rows[0] == ["Word", "Meaning"]
    assert rows[1] == ["ephemeral", "short-lived"]
    assert rows[2] == ["sturdy", "strong"]
    # column order/widths and a frozen header survive the round-trip
    assert ws.column_dimensions["A"].width == round(180 / 7)
    assert ws.freeze_panes == "A2"


# --------------------------------------------------------------------------
# AI answer
# --------------------------------------------------------------------------


def test_ai_answer_needs_key_then_config(client) -> None:
    resp = upload(client, "ai-bank.sat.md", AI_BANK)
    doc_id = resp.json()["id"]
    qid = client.get(f"/api/documents/{doc_id}/questions").json()[0]["id"]

    # key exists in this fixture, but no LLM is configured
    no_config = client.post(
        "/api/ai/answer", json={"document_id": doc_id, "question_id": qid}
    )
    assert no_config.status_code == 409
    assert "Settings" in no_config.json()["detail"]


def test_ai_answer_context_includes_material_options_and_section(
    client, monkeypatch
) -> None:
    resp = upload(client, "ai-bank.sat.md", AI_BANK)
    doc_id = resp.json()["id"]
    qid = client.get(f"/api/documents/{doc_id}/questions").json()[0]["id"]

    captured: dict = {}

    def fake_complete(system: str, messages: list[dict[str, str]]) -> str:
        captured["system"] = system
        captured["messages"] = messages
        return "  The data were found to be reliable because B follows from the premise.  "

    monkeypatch.setattr(llm, "configured", lambda: True)
    monkeypatch.setattr(llm, "complete", fake_complete)

    out = client.post("/api/ai/answer", json={"document_id": doc_id, "question_id": qid})
    assert out.status_code == 200, out.text
    assert out.json()["text"].startswith("The data were")

    user = captured["messages"][0]["content"]
    # Expanded context: material + stem + all options + correct answer + section
    assert "Which choice best completes the sentence?" in user
    assert "Correct answer: B. two" in user
    # Material IS now included
    assert "scholar who counts bees" in user
    # All options ARE now included
    assert "A. one" in user
    assert "B. two" in user
    assert "C. three" in user
    assert "D. four" in user
    # Section info
    assert "Section: RW" in user or "Section: READING & WRITING" in user
    # No hand-written explanation leaks
    assert "private notes" not in user
    assert captured["messages"][0]["role"] == "user"
    assert "SAT tutor" in captured["system"]


def test_ai_answer_without_key_409(client, monkeypatch) -> None:
    doc_id, questions = _mini(client)
    monkeypatch.setattr(llm, "configured", lambda: True)
    out = client.post(
        "/api/ai/answer",
        json={"document_id": doc_id, "question_id": questions[0]["id"]},
    )
    assert out.status_code == 409
    assert "answer key" in out.json()["detail"]


def test_ai_answer_maps_llm_errors(client, monkeypatch) -> None:
    resp = upload(client, "ai-bank.sat.md", AI_BANK)
    doc_id = resp.json()["id"]
    qid = client.get(f"/api/documents/{doc_id}/questions").json()[0]["id"]

    monkeypatch.setattr(llm, "configured", lambda: True)

    def boom(system, messages):
        raise llm.LLMError("authentication failed — check the API key in Settings")

    monkeypatch.setattr(llm, "complete", boom)

    out = client.post("/api/ai/answer", json={"document_id": doc_id, "question_id": qid})
    assert out.status_code == 502
    assert "authentication failed" in out.json()["detail"]
