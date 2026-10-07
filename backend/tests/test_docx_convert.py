"""DOCX -> SAT-MD conversion tests (deterministic, 0 token) plus archive safety."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from app.convert.docx import ConvertError, convert_docx
from app.satmd.parser import parse


def test_converts_sample_docx(sat_docx: Path) -> None:
    result = convert_docx(sat_docx)
    doc = parse(result.satmd)
    assert result.question_count == 5 == len(doc.questions)
    assert result.answers_status == "external"
    assert doc.meta["source"] == "sat-sample.docx"
    assert doc.meta["lang"] == "en"


def test_material_and_stem_split(sat_docx: Path) -> None:
    q1 = parse(convert_docx(sat_docx).satmd).questions[0]
    assert q1.no == 1
    assert q1.material and "reservations" in q1.material
    assert q1.stem.startswith("The author's attitude")
    assert q1.letter_options["B"] == "cautiously optimistic"
    assert q1.answer == "A"


def test_wrapped_option_is_joined(sat_docx: Path) -> None:
    q2 = parse(convert_docx(sat_docx).satmd).questions[1]
    assert q2.letter_options["B"] == (
        "well established and widely cited in the field for nearly a decade"
    )
    assert q2.answer == "C"


def test_image_is_extracted_and_placed(sat_docx: Path) -> None:
    result = convert_docx(sat_docx)
    q3 = parse(result.satmd).questions[2]
    assert "assets/" in q3.stem, q3.stem
    assert q3.images, "image path should be recorded on the question"
    name = q3.images[0].removeprefix("assets/")
    assert name in result.assets
    assert result.assets[name][:8] == b"\x89PNG\r\n\x1a\n"


def test_sections_and_one_line_options(sat_docx: Path) -> None:
    questions = parse(convert_docx(sat_docx).satmd).questions
    assert [q.sec for q in questions[:3]] == ["rw", "rw", "rw"]
    assert [q.sec for q in questions[3:]] == ["math", "math"]
    assert questions[3].letter_options == {"A": "3", "B": "5", "C": "8", "D": "15"}
    assert questions[3].answer == "D"
    assert questions[4].letter_options == {"A": "13", "B": "40", "C": "26", "D": "60"}
    assert questions[4].answer == "A"


def test_answer_key_table_backfills(sat_docx: Path) -> None:
    answers = [q.answer for q in parse(convert_docx(sat_docx).satmd).questions]
    assert answers == ["A", "C", "B", "D", "A"]


def test_answerless_docx_needs_answers(tmp_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("1. Choose the best option.")
    for opt in ["A. one", "B. two", "C. three", "D. four"]:
        doc.add_paragraph(opt)
    path = tmp_path / "no-answers.docx"
    doc.save(path)

    result = convert_docx(path)
    assert result.answers_status == "none"
    assert any("no answers" in w for w in result.warnings)


def test_numberless_docx_is_rejected(tmp_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Which of the following is true?")
    for opt in ["A. one", "B. two", "C. three", "D. four"]:
        doc.add_paragraph(opt)
    path = tmp_path / "numberless.docx"
    doc.save(path)

    with pytest.raises(ConvertError, match="no numbered questions"):
        convert_docx(path)


def test_non_zip_is_rejected(tmp_path: Path) -> None:
    junk = tmp_path / "junk.docx"
    junk.write_bytes(b"this is not a zip container")
    with pytest.raises(ConvertError, match="not a valid"):
        convert_docx(junk)


def test_macro_document_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "macro.docx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("word/vbaProject.bin", b"macro bytes")
    with pytest.raises(ConvertError, match="macro"):
        convert_docx(path)


def test_too_many_entries_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w") as zf:
        for i in range(6000):
            zf.writestr(f"f{i}.bin", b"")
    with pytest.raises(ConvertError, match="too many entries"):
        convert_docx(path)


def test_missing_file_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(ConvertError, match="cannot read"):
        convert_docx(tmp_path / "absent.docx")