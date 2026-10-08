"""Document import, listing, answer-key entry and asset serving."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel

from app import imports as imports_service
from app import llm
from app import repos
from app.db import assets_dir
from app.export import build_export, render_docx_bytes, render_pdf_bytes
from app.repos import words as words_repo
from app.satmd.parser import set_answer
from app.repos import sessions as session_repo

router = APIRouter(prefix="/api/documents", tags=["documents"])

ASSET_NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]*$")
MAX_UPLOAD = imports_service.MAX_UPLOAD
EXPORT_FORMATS = {"pdf", "docx"}

MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class AnswersPayload(BaseModel):
    answers: dict[str, str]


class ExplainPayload(BaseModel):
    content: str = ""


class WordsPayload(BaseModel):
    """Vocabulary grid body: v2 (``columns``/``rows``/``view``) or legacy
    v1 (``headers``/``rows``). Fields are validated in the repo so a single
    ``GridError`` path yields a clean 400."""

    columns: Any = None
    rows: Any = None
    view: Any = None
    headers: Any = None


class NormalizePayload(BaseModel):
    """Normalized question payload returned by /normalize and accepted by /normalize/accept."""
    material: str | None = None
    stem: str
    options: dict[str, str]
    answer: str
    source_ref: str


def _fail(message: str, status: int = 422) -> None:
    raise HTTPException(status_code=status, detail=message)


@router.post("")
async def import_document(file: UploadFile = File(...)) -> dict:
    data = await file.read()
    filename = Path(file.filename or "upload").name
    try:
        imports_service.validate_upload(filename, data)
        result = imports_service.convert_upload(filename, data)
        doc_id = imports_service.commit(result)
    except imports_service.ImportProblem as exc:
        _fail(str(exc), exc.status)
        raise                                          # pragma: no cover

    detail = repos.documents.get_document(doc_id)
    return {**detail, "warnings": result.warnings}


@router.get("")
def list_documents() -> list[dict]:
    return repos.documents.list_documents()


@router.get("/{doc_id}")
def get_document(doc_id: int) -> dict:
    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    return doc


@router.get("/{doc_id}/questions")
def document_questions(doc_id: int) -> list[dict]:
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    return repos.documents.list_questions(doc_id)


@router.get("/{doc_id}/history")
def document_history(doc_id: int) -> list[dict]:
    return session_repo.history(doc_id)


@router.put("/{doc_id}/answers")
def put_answers(doc_id: int, payload: AnswersPayload) -> dict:
    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover

    cleaned: dict[str, str] = {}
    for ext_id, letter in payload.answers.items():
        letter = (letter or "").strip().upper()
        if letter not in {"A", "B", "C", "D"}:
            _fail(f"invalid answer {letter!r} for {ext_id} (expected A-D)", 400)
        cleaned[ext_id] = letter
    if not cleaned:
        _fail("no answers supplied", 400)

    changed = repos.documents.set_answers(doc_id, cleaned)
    if changed == 0:
        _fail("no matching questions", 404)

    # keep the on-disk SAT-MD in sync (source of truth for export/debug)
    try:
        text = repos.documents.read_satmd(doc_id)
        for ext_id, letter in cleaned.items():
            text = set_answer(text, ext_id, letter)
        text = re.sub(r"(?m)^answers:\s*\S+\s*$", "answers: external", text, count=1)
        repos.documents.write_satmd(doc_id, text)
    except (KeyError, OSError):
        pass                                            # DB already updated

    return repos.documents.get_document(doc_id)


@router.put("/{doc_id}/questions/{question_id}/explain")
def put_explain(doc_id: int, question_id: int, payload: ExplainPayload) -> dict:
    """Hand-written explanation — saved straight to the DB (0 token)."""
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    content = payload.content.strip() or None
    if not repos.documents.set_explain(doc_id, question_id, content):
        _fail("question not found", 404)
        raise                                          # pragma: no cover
    return {"question_id": question_id, "explain": content}


# --------------------------------------------------------------------------
# CB-style normalization (batch 12)
# --------------------------------------------------------------------------


@router.post("/{doc_id}/questions/{question_id}/normalize")
def normalize_question(doc_id: int, question_id: int) -> dict:
    """Generate a CB-style normalization for a question (LLM call; explicit user action)."""
    if not llm.configured():
        _fail("no LLM configured — set API key and model in Settings", 409)
        raise                                          # pragma: no cover

    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover

    questions = repos.documents.list_questions(doc_id)
    q = next((qq for qq in questions if qq["id"] == question_id), None)
    if q is None:
        _fail("question not found", 404)
        raise                                          # pragma: no cover

    # Need material, stem, options, answer, source_ref
    # list_questions returns the right shape.
    material = q.get("material")
    stem = q.get("stem", "")
    options = q.get("options", {})
    answer = q.get("answer", "")
    source_ref = q.get("source", "")

    if not stem or not options or answer not in {"A", "B", "C", "D"}:
        _fail("question missing required fields for normalization", 400)
        raise                                          # pragma: no cover

    try:
        from app.llm import normalize as normalize_mod

        result = normalize_mod.normalize_question(
            material=material,
            stem=stem,
            options=options,
            answer=answer,
            source_ref=source_ref,
        )
    except normalize_mod.NormalizeError as exc:
        _fail(str(exc), 422)
        raise                                          # pragma: no cover
    except llm.LLMError as exc:
        _fail(f"LLM request failed: {exc}", 502)
        raise                                          # pragma: no cover

    return {
        "original": result.original,
        "normalized": result.normalized,
        "changed": result.changed,
        "answer_preserved": result.answer_preserved,
        "requires_review": result.requires_review,
    }


@router.post("/{doc_id}/questions/{question_id}/normalize/accept")
def accept_normalization(doc_id: int, question_id: int, payload: NormalizePayload) -> dict:
    """Persist an accepted normalization (validates invariants again)."""
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover

    try:
        from app.llm import normalize as normalize_mod

        return normalize_mod.accept_normalization(doc_id, question_id, payload.model_dump())
    except normalize_mod.NormalizeError as exc:
        _fail(str(exc), 422)
        raise                                          # pragma: no cover


@router.get("/{doc_id}/words")
def get_words(doc_id: int) -> dict:
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    return words_repo.get_grid(doc_id)


@router.put("/{doc_id}/words")
def put_words(doc_id: int, payload: WordsPayload) -> dict:
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    try:
        return words_repo.save_grid(doc_id, payload.model_dump(exclude_none=True))
    except words_repo.GridError as exc:
        _fail(str(exc), 400)
        raise                                          # pragma: no cover


@router.get("/{doc_id}/words/export")
def export_words(doc_id: int) -> Response:
    """Vocabulary grid as .xlsx (openpyxl; Numbers opens it too)."""
    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    grid = words_repo.get_grid(doc_id)
    columns = grid["columns"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Vocabulary"
    ws.append([c["name"] for c in columns])
    for row in grid["rows"]:
        ws.append([row["cells"].get(c["id"], "") for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, start=1):
        # xlsx width is in characters; px→chars ≈ width / 7, clamped.
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(60, round(col["width"] / 7)))

    buf = BytesIO()
    wb.save(buf)

    filename = f"{_slugify(doc['title'])}-vocabulary.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _slugify(title: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", title).strip("-.")[:60]
    return slug or "document"


@router.get("/{doc_id}/export/{fmt}")
def export_document(doc_id: int, fmt: str) -> Response:
    """Export questions + answers + explanations as PDF or DOCX (0 token)."""
    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    if fmt not in EXPORT_FORMATS:
        _fail("format must be pdf or docx", 400)
        raise                                          # pragma: no cover

    export_doc = build_export(doc_id)
    if export_doc is None:                             # pragma: no cover
        _fail("document not found", 404)
        raise
    content = render_pdf_bytes(export_doc) if fmt == "pdf" else render_docx_bytes(export_doc)
    return Response(
        content=content,
        media_type=MEDIA_TYPES[fmt],
        headers={
            "Content-Disposition": f'attachment; filename="{_slugify(doc["title"])}.{fmt}"'
        },
    )


@router.delete("/{doc_id}")
def delete_document(doc_id: int) -> dict:
    if not repos.documents.delete_document(doc_id):
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    return {"deleted": doc_id}


@router.get("/{doc_id}/assets/{name}")
def document_asset(doc_id: int, name: str) -> FileResponse:
    if not ASSET_NAME_RE.match(name) or ".." in name:
        _fail("invalid asset name", 400)
        raise                                          # pragma: no cover
    path = assets_dir(doc_id) / name
    if not path.is_file():
        _fail("asset not found", 404)
        raise                                          # pragma: no cover
    media = "image/png" if name.lower().endswith(".png") else "application/octet-stream"
    return FileResponse(path, media_type=media)
