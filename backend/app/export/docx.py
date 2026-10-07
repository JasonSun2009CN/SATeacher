"""DOCX export via python-docx (deterministic, 0 token)."""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from app.export import math_render
from app.export.model import ExportDoc, section_label

_SIZE_RE = re.compile(r'width="([\d.]+)ex"\s+height="([\d.]+)ex"')
_PT_PER_EX = 4.6
_MAX_PT = 440.0


def _svg_size(svg: str) -> tuple[float, float]:
    match = _SIZE_RE.search(svg)
    if not match:
        return (0.0, 0.0)
    return (float(match.group(1)), float(match.group(2)))


def _svg_to_png(svg: str) -> tuple[bytes | None, float]:
    try:
        import cairosvg

        png = cairosvg.svg2png(bytestring=svg.encode("utf-8"), scale=3)
    except Exception:
        return (None, 0.0)
    width_ex, _ = _svg_size(svg)
    width_pt = min(width_ex * _PT_PER_EX, _MAX_PT) if width_ex else 40.0
    return (png, max(width_pt, 6.0))


def _add_rich(document: Document, text: str, svgs: dict, **para_kwargs):
    para = document.add_paragraph(**para_kwargs)
    for kind, value in math_render.split_math(text or ""):
        if kind == "text":
            lines = value.split("\n")
            for i, line in enumerate(lines):
                if i:
                    para.add_run().add_break()
                if line:
                    para.add_run(line)
            continue
        tex, _display = value
        svg = svgs.get(value)
        png, width_pt = _svg_to_png(svg) if svg else (None, 0.0)
        if png:
            para.add_run().add_picture(BytesIO(png), width=Pt(width_pt))
        else:
            para.add_run(f"${tex}$")
    return para


def _all_text(doc: ExportDoc) -> list[str]:
    texts = [doc.title, doc.source_filename, *doc.vocab_headers]
    for q in doc.questions:
        texts += [q.material or "", q.stem, q.explain or "", *q.options.values()]
    for row in doc.vocab_rows:
        texts.extend(row)
    return texts


def render_docx(doc: ExportDoc) -> bytes:
    svgs = math_render.render(math_render.collect_math(_all_text(doc)))
    document = Document()

    document.add_heading(doc.title, level=0)
    meta = document.add_paragraph()
    run = meta.add_run(
        f"{doc.source_filename} · {len(doc.questions)} questions · {doc.answered} answered"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    document.add_heading("Questions", level=1)
    for q in doc.questions:
        document.add_heading(f"Question {q.no} · {section_label(q.sec)}", level=2)
        if q.source:
            src = document.add_paragraph()
            r = src.add_run(q.source)
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        if q.material:
            _add_rich(document, q.material, svgs, style="Intense Quote")
        _add_rich(document, q.stem, svgs)
        for letter in "ABCD":
            if letter in q.options:
                _add_rich(document, f"{letter}. {q.options[letter]}", svgs, style="List Bullet")
        for image in q.images:
            try:
                document.add_picture(str(image), width=Inches(5.5))
            except Exception:
                pass

    document.add_heading("Answer Key", level=1)
    key = document.add_paragraph()
    if any(q.answer for q in doc.questions):
        key.add_run(
            "   ".join(f"{q.no}. {q.answer or '—'}" for q in doc.questions)
        )
    else:
        key.add_run("No answers.").font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    document.add_heading("Explanations", level=1)
    explained = [q for q in doc.questions if q.explain]
    if not explained:
        r = document.add_paragraph().add_run("No explanations saved yet.")
        r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    for q in explained:
        para = _add_rich(document, q.explain, svgs)
        para.runs[0].text = f"Question {q.no}. " + para.runs[0].text

    if doc.vocab_rows:
        document.add_heading("Vocabulary", level=1)
        table = document.add_table(rows=1, cols=len(doc.vocab_headers))
        table.style = "Table Grid"
        for i, header in enumerate(doc.vocab_headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(header)
            run.bold = True
        for row in doc.vocab_rows:
            cells = table.add_row().cells
            for i, cell in enumerate(row):
                if i < len(cells):
                    cells[i].text = cell

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()